"""Write OCP field-level mapping workbook from a project JSON file."""

from __future__ import annotations

import copy
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from excel_to_json_mapping import STAGE_SHEETS, _is_placeholder_data_element

# 1-based column indices on OCDS stage sheets (mapping, example, notes)
MAPPING_COL = 6
EXAMPLE_COL = 7
NOTES_COL = 8
PATH_COL = 3

FIELD_ROW_TYPES = {
    "required_field",
    "field",
    "additional_field",
    "extension_field",
}

DATA_SOURCE_FIELDS = [
    "short_name",
    "data_source",
    "vendor",
    "url_documentation",
    "agency_unit_responsible",
    "planning",
    "tender",
    "award",
    "contract",
    "implementation_physical",
    "implementation_financial",
    "revision_history",
    "documents_available",
    "existing_apis_data_publication",
]

DATA_ELEMENT_BASE_FIELDS = [
    "for_mapping",
    "data_source",
    "section_table",
    "data_element",
    "publish",
    "example",
    "description",
    "data_type",
    "notes",
    "mapped",
    "mapping_details",
]


def _cell_str(val: Any) -> str:
    if val is None:
        return ""
    return str(val).strip()


def _find_header_row(ws, header_label: str, col: int = 1) -> int | None:
    target = header_label.strip().lower()
    for r in range(1, ws.max_row + 1):
        if _cell_str(ws.cell(r, col).value).lower() == target:
            return r
    return None


def _delete_rows_from(ws, start_row: int, predicate) -> int:
    removed = 0
    r = start_row
    while r <= ws.max_row:
        if predicate(r):
            ws.delete_rows(r, 1)
            removed += 1
        else:
            r += 1
    return removed


def write_data_sources_sheet(ws, project: dict) -> int:
    header_row = _find_header_row(ws, "short name")
    if header_row is None:
        return 0
    _delete_rows_from(
        ws,
        header_row + 1,
        lambda r: any(_cell_str(ws.cell(r, c).value) for c in range(1, 15)),
    )
    rows = project.get("data_sources", {}).get("rows") or []
    r = header_row + 1
    for row in rows:
        for j, key in enumerate(DATA_SOURCE_FIELDS, start=1):
            ws.cell(r, j).value = row.get(key, "") or None
        r += 1
    return len(rows)


def write_data_elements_sheet(ws, project: dict) -> int:
    header_row = _find_header_row(ws, "for mapping")
    if header_row is None:
        return 0
    _delete_rows_from(
        ws,
        header_row + 1,
        lambda r: any(_cell_str(ws.cell(r, c).value) for c in range(1, 12)),
    )
    de = project.get("data_elements") or {}
    columns = de.get("columns") or []
    col_keys = [c.get("key", "") for c in columns] if columns else DATA_ELEMENT_BASE_FIELDS
    rows_out = []
    for row in de.get("rows") or []:
        if _is_placeholder_data_element(
            row.get("for_mapping", ""),
            row.get("section_table", ""),
            row.get("data_element", ""),
        ):
            continue
        rows_out.append(row)
    r = header_row + 1
    for row in rows_out:
        for j, key in enumerate(col_keys, start=1):
            val = row.get(key, "")
            if key.startswith("count_") or key.startswith("presence_"):
                ws.cell(r, j).value = int(val) if val not in ("", None) else 0
            else:
                ws.cell(r, j).value = val if val != "" else None
        r += 1
    return len(rows_out)


def apply_mapping_edits(wb, mapping_edits: dict[str, dict]) -> int:
    """Write tier-C edits onto OCDS stage sheets. Returns cells written."""
    written = 0
    for sheet_name, stage_key in STAGE_SHEETS.items():
        if sheet_name not in wb.sheetnames:
            continue
        ws = wb[sheet_name]
        for r in range(1, ws.max_row + 1):
            rtype = _cell_str(ws.cell(r, 1).value).lower()
            if rtype not in FIELD_ROW_TYPES:
                continue
            path = _cell_str(ws.cell(r, PATH_COL).value)
            if not path:
                continue
            edit = mapping_edits.get(f"{stage_key}::{path}")
            if not edit:
                continue
            for col, key in ((MAPPING_COL, "mapping"), (EXAMPLE_COL, "example"), (NOTES_COL, "notes")):
                if key in edit and edit[key] not in (None, ""):
                    ws.cell(r, col).value = edit[key]
                    written += 1
    return written


def export_ocp_workbook(project: dict, template: Path, output: Path) -> dict[str, int]:
    if not template.exists():
        raise FileNotFoundError(f"Template workbook not found: {template}")
    wb = load_workbook(template)
    stats = {
        "data_sources_written": 0,
        "data_elements_written": 0,
        "mapping_cells_written": 0,
    }
    if "1. Data Sources" in wb.sheetnames:
        stats["data_sources_written"] = write_data_sources_sheet(wb["1. Data Sources"], project)
    if "2. Data Elements" in wb.sheetnames:
        stats["data_elements_written"] = write_data_elements_sheet(wb["2. Data Elements"], project)
    stats["mapping_cells_written"] = apply_mapping_edits(wb, project.get("mapping_edits") or {})
    output.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output)
    return stats


def build_mapping_yaml_document(project: dict) -> dict:
    """Build a YAML-serializable document for ETL pipelines."""
    meta = copy.deepcopy(project.get("metadata") or {})
    mapping_edits = project.get("mapping_edits") or {}
    by_stage: dict[str, dict[str, dict]] = {}
    for key, edit in mapping_edits.items():
        if "::" not in key:
            continue
        stage, path = key.split("::", 1)
        if not path or not edit.get("mapping"):
            continue
        by_stage.setdefault(stage, {})[path] = {
            "for_mapping": edit.get("mapping", ""),
            **({"example": edit["example"]} if edit.get("example") else {}),
            **({"notes": edit["notes"]} if edit.get("notes") else {}),
        }

    elements = []
    for row in project.get("data_elements", {}).get("rows") or []:
        if _is_placeholder_data_element(
            row.get("for_mapping", ""),
            row.get("section_table", ""),
            row.get("data_element", ""),
        ):
            continue
        elements.append(
            {
                "key": row.get("key", ""),
                "for_mapping": row.get("for_mapping", ""),
                "data_source": row.get("data_source", ""),
                "section_table": row.get("section_table", ""),
                "data_element": row.get("data_element", ""),
                **({"mapping_details": row["mapping_details"]} if row.get("mapping_details") else {}),
                **({"publish": row["publish"]} if row.get("publish") else {}),
            }
        )

    proposals = []
    for p in project.get("extension_proposals") or []:
        proposals.append(
            {
                "extension": p.get("extension", ""),
                "path": p.get("path", ""),
                "title": p.get("title", ""),
                "data_type": p.get("data_type", ""),
                "status": p.get("status", "draft"),
                **({"for_mapping": p["for_mapping"]} if p.get("for_mapping") else {}),
                **({"description": p["description"]} if p.get("description") else {}),
            }
        )

    doc = {
        "version": 1,
        "meta": {
            "project_id": meta.get("project_id", ""),
            "name": meta.get("name", ""),
            "source_system": meta.get("source_system", ""),
            "ocds_version": meta.get("ocds_version", "1.1.5"),
            "exported_at": meta.get("updated_at", ""),
        },
        "data_elements": elements,
        "ocds_mappings": by_stage,
        **({"extension_proposals": proposals} if proposals else {}),
    }
    pub = _publication_for_export(project.get("publication_config"))
    if pub:
        doc["publication"] = pub
    return doc


def _publication_for_export(raw: dict | None) -> dict | None:
    if not raw:
        return None
    out: dict = {}
    if raw.get("version"):
        out["version"] = raw["version"]
    if raw.get("ocid_prefix"):
        out["ocid_prefix"] = raw["ocid_prefix"]
    pub = raw.get("publisher") or {}
    if any(pub.get(k) for k in ("name", "scheme", "uid", "uri")):
        out["publisher"] = {k: pub[k] for k in ("name", "scheme", "uid", "uri") if pub.get(k)}
    if raw.get("license"):
        out["license"] = raw["license"]
    if raw.get("publicationPolicy"):
        out["publicationPolicy"] = raw["publicationPolicy"]
    if raw.get("extensions"):
        out["extensions"] = raw["extensions"]
    defs = raw.get("defaults") or {}
    if defs:
        out["defaults"] = {k: v for k, v in defs.items() if v}
    gen = raw.get("generated") or {}
    gen_out: dict = {}
    if gen.get("ocid", {}).get("from"):
        gen_out["ocid"] = {"from": gen["ocid"]["from"]}
    if gen.get("release_id"):
        gen_out["release_id"] = gen["release_id"]
    if gen.get("release_date", {}).get("prefer"):
        gen_out["release_date"] = {"prefer": gen["release_date"]["prefer"]}
    if gen_out:
        out["generated"] = gen_out
    if raw.get("planning_triggers"):
        out["planning_triggers"] = raw["planning_triggers"]
    buyer = raw.get("buyer") or {}
    if buyer.get("name") or buyer.get("id"):
        out["buyer"] = {k: buyer[k] for k in ("name", "id") if buyer.get(k)}
    return out if len(out) > 1 else None
