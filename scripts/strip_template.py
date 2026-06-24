#!/usr/bin/env python3
"""Create a blank OCP field-level mapping workbook from a filled template.

Clears user-entered data while preserving structure, formulas, and formatting:
  - Sheets 1 & 2: data rows (keeps header + one optional example row)
  - OCDS stage sheets: mapping (col F), example (G), notes (H) — 1-based cols 6–8

Requires openpyxl.

Usage:
  python scripts/strip_template.py
  python scripts/strip_template.py -i references/filled.xlsx -o references/blank.xlsx
"""

from __future__ import annotations

import argparse
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
REFERENCES = ROOT / "references"
DEFAULT_INPUT = REFERENCES / "OCDS_Mapping_Template_1.1.xlsx"
DEFAULT_OUTPUT = REFERENCES / "OCDS_Mapping_Template_1.1_blank.xlsx"

STAGE_SHEETS = [
    "(OCDS) 1. General (all stages)",
    "(OCDS) 2. Planning",
    "(OCDS) 3. Tender",
    "(OCDS) 4. Award",
    "(OCDS) 5. Contract",
    "(OCDS) 6. Implementation",
]

# 1-based column indices on OCDS stage sheets (mapping, example, notes)
MAPPING_COLS = (6, 7, 8)


def _cell_str(val) -> str:
    if val is None:
        return ""
    return str(val).strip()


def clear_data_sources_sheet(ws, keep_example: bool = False) -> int:
    """Remove data rows below header; optionally keep first example row."""
    header_row = None
    for r in range(1, ws.max_row + 1):
        if _cell_str(ws.cell(r, 1).value).lower() == "short name":
            header_row = r
            break
    if header_row is None:
        return 0
    removed = 0
    r = header_row + 1
    while r <= ws.max_row:
        if not any(_cell_str(ws.cell(r, c).value) for c in range(1, 8)):
            r += 1
            continue
        short = _cell_str(ws.cell(r, 1).value).lower()
        if keep_example and short == "example":
            r += 1
            continue
        ws.delete_rows(r, 1)
        removed += 1
    return removed


def clear_data_elements_sheet(ws, keep_example: bool = False) -> int:
    header_row = None
    for r in range(1, ws.max_row + 1):
        if _cell_str(ws.cell(r, 1).value).lower() == "for mapping":
            header_row = r
            break
    if header_row is None:
        return 0
    removed = 0
    r = header_row + 1
    while r <= ws.max_row:
        if not any(_cell_str(ws.cell(r, c).value) for c in range(1, 12)):
            r += 1
            continue
        first = _cell_str(ws.cell(r, 1).value)
        field = _cell_str(ws.cell(r, 4).value)
        if keep_example and first.lower() == "example":
            r += 1
            continue
        if first.lower() in ("for mapping", "") and not field:
            r += 1
            continue
        ws.delete_rows(r, 1)
        removed += 1
    return removed


def clear_stage_sheet(ws) -> int:
    """Clear mapping/example/notes on rows whose col A is a field row type."""
    field_types = {
        "required_field",
        "field",
        "additional_field",
        "extension_field",
    }
    cleared = 0
    for r in range(1, ws.max_row + 1):
        rtype = _cell_str(ws.cell(r, 1).value).lower()
        if rtype not in field_types:
            continue
        for c in MAPPING_COLS:
            if ws.cell(r, c).value not in (None, ""):
                ws.cell(r, c).value = None
                cleared += 1
    return cleared


def strip_workbook(path_in: Path, path_out: Path, keep_example: bool = False) -> dict:
    wb = load_workbook(path_in)
    stats = {"data_sources_rows_removed": 0, "data_elements_rows_removed": 0, "stage_cells_cleared": 0}

    if "1. Data Sources" in wb.sheetnames:
        stats["data_sources_rows_removed"] = clear_data_sources_sheet(
            wb["1. Data Sources"], keep_example=keep_example
        )
    if "2. Data Elements" in wb.sheetnames:
        stats["data_elements_rows_removed"] = clear_data_elements_sheet(
            wb["2. Data Elements"], keep_example=keep_example
        )
    for name in STAGE_SHEETS:
        if name in wb.sheetnames:
            stats["stage_cells_cleared"] += clear_stage_sheet(wb[name])

    path_out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path_out)
    return stats


def main() -> int:
    p = argparse.ArgumentParser(description=__doc__)
    p.add_argument("--input", "-i", type=Path, default=DEFAULT_INPUT)
    p.add_argument("--output", "-o", type=Path, default=DEFAULT_OUTPUT)
    p.add_argument("--keep-example", action="store_true", help="Keep Example rows on sheets 1–2")
    args = p.parse_args()

    if not args.input.exists():
        raise SystemExit(f"Input not found: {args.input}")

    print(f"Stripping {args.input.name} -> {args.output.name}")
    stats = strip_workbook(args.input, args.output, keep_example=args.keep_example)
    for k, v in stats.items():
        print(f"  {k}: {v}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
